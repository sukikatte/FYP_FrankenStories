#!/usr/bin/env python3
# -*- coding: utf-8 -*-


from __future__ import annotations

import argparse
import copy
import csv
import json
import sys
import time
from pathlib import Path
from typing import Any, Dict, List

import openai


EXP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = EXP_DIR.parent

# CONFIG — paste keys and model ids here
OPENAI_API_KEY = ""
TOGETHER_API_KEY = ""
OPENAI_BASE_URL = ""
TOGETHER_BASE_URL = "https://api.together.xyz/v1"

SPEC_STORY_ELEMENTS_DIR = PROJECT_ROOT / "spec_library"
TEMPLATE_FILE = PROJECT_ROOT.parent.parent / "template.xml"
OUTPUT_BASE = PROJECT_ROOT / "outputs_insights" / "time_1"
_XLSX_DEFAULT = PROJECT_ROOT / "pose database.xlsx"
DEFAULT_POSE_DB = _XLSX_DEFAULT if _XLSX_DEFAULT.exists() else (EXP_DIR / "pose_database.json")

SLEEP_BETWEEN_CALLS = 1.0
DEFAULT_POSE_FALLBACK_DESCRIPTION = (
    "No separate gloss in pose_database; interpret staging from the canonical pose name."
)

DIMENSIONS: Dict[str, Dict] = {
    "dim1_single_char_single_story": {"char_count": 1, "story_count": "single"},
    "dim2_single_char_multi_story": {"char_count": 1, "story_count": "multi"},
    "dim3_multi_char_single_story": {"char_count": 2, "story_count": "single"},
    "dim4_multi_char_multi_story": {"char_count": 2, "story_count": "multi"},
}

OPENAI_MODEL_GPT41 = "gpt-4.1"
OPENAI_MODEL_GPT4O = "gpt-4o"
TOGETHER_MODEL_DEEPSEEK_R1 = "deepseek-ai/DeepSeek-R1"
TOGETHER_MODEL_DEEPSEEK_V3 = "deepseek-ai/DeepSeek-V3.1"
TOGETHER_MODEL_LLAMA = "meta-llama/Llama-3.3-70B-Instruct-Turbo"

MODELS = {
    "gpt-4.1": ("gpt_4.1", "CoT", OPENAI_MODEL_GPT41),
    "gpt-4o": ("gpt_4o", "CoT", OPENAI_MODEL_GPT4O),
    "DeepSeek-R1": ("deepseek_r1", "CoT", TOGETHER_MODEL_DEEPSEEK_R1),
    "DeepSeek-V3": ("deepseek_v3", "CoT", TOGETHER_MODEL_DEEPSEEK_V3),
    "llama-3.3-70B": ("llama_3_70b", "Direct", TOGETHER_MODEL_LLAMA),
}

TOGETHER_API_MODEL_IDS = frozenset(
    {
        TOGETHER_MODEL_DEEPSEEK_R1,
        TOGETHER_MODEL_DEEPSEEK_V3,
        TOGETHER_MODEL_LLAMA,
    }
)

FIVE_MODEL_DISPLAY_NAMES = [
    "gpt-4.1",
    "gpt-4o",
    "DeepSeek-V3",
    "DeepSeek-R1",
    "llama-3.3-70B",
]


def load_pose_database_xlsx(path: Path) -> Dict[str, str]:
    """Load pose -> description from Excel (e.g. columns Pose, Type, Depicts)."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError(
            "Reading .xlsx requires openpyxl. Install with: pip install openpyxl"
        ) from e

    mapping: Dict[str, str] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(min_row=1, values_only=True)
        header = next(rows, None)
        if not header:
            return {}
        names = [
            (str(c).strip() if c is not None else "").lower() for c in header
        ]

        def col_idx(*candidates: str) -> int | None:
            for i, n in enumerate(names):
                if n in candidates:
                    return i
            return None

        i_pose = col_idx("pose")
        if i_pose is None:
            raise ValueError(
                f"Excel {path} needs a 'Pose' column; got headers {list(header)}"
            )
        i_type = col_idx("type")
        i_depicts = col_idx("depicts", "description", "desc")

        def cell(row: tuple, idx: int | None) -> str:
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            if v is None:
                return ""
            return str(v).strip()

        for row in rows:
            if row is None:
                continue
            pose = cell(row, i_pose)
            if not pose:
                continue
            parts: List[str] = []
            dep = cell(row, i_depicts)
            typ = cell(row, i_type)
            if dep:
                parts.append(dep)
            if typ:
                parts.append(f"type: {typ}")
            mapping[pose] = ". ".join(parts) if parts else ""
    finally:
        wb.close()

    return mapping


def load_pose_database(path: Path) -> Dict[str, str]:
    """Load pose -> description. Supports .xlsx, .csv, or .json."""
    if not path.exists():
        return {}
    suf = path.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        return load_pose_database_xlsx(path)
    if suf == ".csv":
        mapping: Dict[str, str] = {}
        with path.open("r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                return {}
            pkey = None
            dkey = None
            for h in reader.fieldnames:
                hl = h.strip().lower()
                if hl in ("pose", "pose_name", "name", "id"):
                    pkey = h
                if hl in ("description", "desc", "gloss", "text"):
                    dkey = h
            if not pkey or not dkey:
                raise ValueError(
                    f"CSV {path} needs recognizable pose and description columns "
                    f"(e.g. pose,description); got {reader.fieldnames}"
                )
            for row in reader:
                k = (row.get(pkey) or "").strip()
                v = (row.get(dkey) or "").strip()
                if k:
                    mapping[k] = v
        return mapping
    if suf != ".json":
        raise ValueError(f"Unsupported pose database format: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict) and all(isinstance(k, str) for k in data):
        # Allow {"poses": {...}} wrapper
        if "poses" in data and isinstance(data["poses"], dict):
            raw = data["poses"]
        else:
            raw = data
        return {str(k): str(v) if v is not None else "" for k, v in raw.items()}
    raise ValueError(f"Unsupported pose database format: {path}")


def annotate_elements_for_prompt(elements: Dict[str, Any], pose_db: Dict[str, str]) -> Dict[str, Any]:
    """Deep copy; replace poses list with 'name (description)' strings for the prompt only."""
    out = copy.deepcopy(elements)
    poses = out.get("poses", [])
    if not isinstance(poses, list):
        return out
    new_poses: List[str] = []
    for p in poses:
        if not isinstance(p, str):
            continue
        desc = (pose_db.get(p) or "").strip()
        if not desc:
            desc = DEFAULT_POSE_FALLBACK_DESCRIPTION
        new_poses.append(f"{p} ({desc})")
    out["poses"] = new_poses
    return out


def load_template() -> str:
    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError(f"template.xml not found at {TEMPLATE_FILE}")
    return TEMPLATE_FILE.read_text(encoding="utf-8")


def load_story_elements(story_id: int) -> Dict:
    story_file = SPEC_STORY_ELEMENTS_DIR / f"story_{story_id}.json"
    if not story_file.exists():
        raise FileNotFoundError(f"Story elements file not found: {story_file}")
    with story_file.open("r", encoding="utf-8") as f:
        return json.load(f)


def merge_story_elements(story_ids: List[int]) -> Dict:
    merged = {
        "characters": {},
        "poses": set(),
        "settings": set(),
        "facings": set(),
        "positions": set(),
        "balloon_status": set(),
    }

    for sid in story_ids:
        data = load_story_elements(sid)

        for char_id, char_data in data.get("characters", {}).items():
            if char_id not in merged["characters"]:
                merged["characters"][char_id] = char_data
            else:
                existing_variants = merged["characters"][char_id].get("variants", {})
                new_variants = char_data.get("variants", {})
                existing_variants.update(new_variants)
                merged["characters"][char_id]["variants"] = existing_variants

        merged["poses"].update(data.get("poses", []))
        merged["settings"].update(data.get("settings", []))
        merged["facings"].update(data.get("facings", []))
        merged["positions"].update(data.get("positions", []))
        merged["balloon_status"].update(data.get("balloon_status", []))

    return {
        "characters": merged["characters"],
        "poses": sorted(merged["poses"]),
        "settings": sorted(merged["settings"]),
        "facings": sorted(merged["facings"]),
        "positions": sorted(merged["positions"]),
        "balloon_status": sorted(merged["balloon_status"]),
    }


def chat_model(
    client: openai.OpenAI,
    model_name: str,
    messages,
    max_tokens: int = 4000,
    temperature: float = 0.4,
) -> str:
    params = {"model": model_name, "messages": messages, "temperature": temperature}
    if model_name.startswith("o1"):
        params["max_completion_tokens"] = max_tokens
    else:
        params["max_tokens"] = max_tokens

    try:
        resp = client.chat.completions.create(**params)
        content = resp.choices[0].message.content
        if content:
            return content
        raise RuntimeError("API returned empty content.")
    except Exception as e:
        raise RuntimeError(f"API call failed: {e}") from e


def pick_client(
    api_model: str,
    openai_client: openai.OpenAI,
    together_client: openai.OpenAI,
) -> openai.OpenAI:
    if api_model in TOGETHER_API_MODEL_IDS:
        return together_client
    return openai_client


def build_user_prompt(
    dimension_id: str,
    story_label: str,
    elements_for_prompt: Dict,
    template_xml: str,
    char_count: int,
    story_count: str,
) -> str:
    dim_desc = {
        "dim1_single_char_single_story": "Dim1: Single character and single story.",
        "dim2_single_char_multi_story": "Dim2: Single character and multiple stories merged.",
        "dim3_multi_char_single_story": "Dim3: Multiple characters within a single story.",
        "dim4_multi_char_multi_story": "Dim4: Multiple characters and multiple stories merged.",
    }.get(dimension_id, dimension_id)

    if char_count == 1:
        char_constraint = (
            "- You MUST use only ONE main character ID from the element library\n"
            "  across the whole story (single-character setting).\n"
        )
    else:
        char_constraint = (
            "- You MUST use MULTIPLE distinct character IDs from the element library\n"
            "  across the story (multi-character setting), and most panels should\n"
            "  contain 2 characters.\n"
        )

    story_info = (
        f"You must create a ComicXML story for **{dim_desc}**\n"
        f"- Story source type: {story_count} (story_label = {story_label})\n"
        "- Panel count must be between 6 and 12 (inclusive).\n"
        f"{char_constraint}"
    )

    elements_json = json.dumps(elements_for_prompt, ensure_ascii=False, indent=2)

    prompt = f"""
You are an expert ComicXML generation assistant.

Your task:
- Generate a COMPLETE ComicXML for a comic story.
- STRICTLY follow the structure defined by the template.
- Use ONLY elements from the provided element library.
- Respect the character constraints above (single-character or multi-character).

{story_info}

Element library (permitted values only). Poses are listed as **canonical_pose_name (description)**.
The parenthetical text is ONLY a staging gloss for you; in the final XML you must still obey
the canonical element rules below.

{elements_json}

ComicXML template (STRUCTURE GUIDE, not actual values):
{template_xml}

IMPORTANT CONSTRAINTS:
1. XML must be well-formed (proper tag closure and nesting).
2. XML must follow the template structure (no missing required elements).
3. Use only elements (characters, poses, settings, positions, facings, balloon_status) from
   the element library above. For **poses**, each entry is `CANONICAL_NAME (description)`.
   The canonical name may contain spaces; it is the substring **before the first ` (`**
   (space + opening parenthesis). Put **only** that canonical name inside `<pose>...</pose>`,
   never the parenthetical description.
4. You MAY reuse elements across panels, but do not invent new ones.
5. Panel count must be between 6 and 12.
6. For multi-character dimensions, most panels should contain multiple characters.
7. Pay close attention to facing directions in multi-character panels.
8. If two characters are in a face-to-face conversation across left and right positions, the usual orientation is:
   - left-position character: facing="right"
   - right-position character: facing="left"
9. In non-conversational or non-face-to-face situations, choose facing directions according to the narrative context of that panel.
10. **Cast visibility across panels (multi-character dimensions):**
   - Default: If panel N and panel N+1 depict the **same ongoing scene** and the **same interaction or conversation** continues, keep the **same characters visibly on stage** in panel N+1 unless the story motivates a change.
   - Do **not** drop a co-present character between adjacent panels **without a clear diegetic reason** (e.g. they exit on-panel, the `<below>` caption states a time jump or location change, or the dialogue explicitly addresses their absence).
   - A **single-character panel** is allowed when: (a) the story **explicitly** moves to a new beat with only one character; (b) a `<below>` caption states a transition; or (c) the plot states the other character left or is off-screen for a named reason.
   - After **physical actions involving two characters** (rescuing, dousing, grabbing, face-to-face dialogue), the **next panel should normally still show both**, repositioned if needed, rather than removing one character for a solo middle composition without motivation.
11. **Layout:** When two characters are **co-present in the same narrative beat**, prefer **left+right** (or **left+middle+right**) over a **solo `<middle>`** panel. Do not use solo middle only to emphasize a speech balloon if both characters should still be in the scene.
12. **`<below>` captions:** For panels where **multiple characters share a beat** (especially immediately after joint action or dialogue), write each relevant `<below>` so it **mentions where each of them is** (stance or position relative to the other), not only one of them — unless a transition or solo beat is intentional and stated in that caption.
13. **Global story logic and continuity:** The whole comic must read as one coherent story arc (setup -> development -> consequence/resolution), not disconnected moments. Maintain clear cause-effect links across adjacent panels. If the story changes location, time, or cast composition, explicitly signal that change in `<below>` and/or dialogue.
14. **Element whitelist (strict):** Do not output any out-of-library value, including near-synonyms or morphology variants (e.g. `laughing` when only `reacting` exists). If an intended action is unavailable, pick the closest allowed element from the provided library.
15. **Cross-panel position consistency (multi-character dimensions):** If the same two characters continue the same interaction across adjacent panels in the same scene, keep their screen sides stable by default (the character on `left` stays on `left`, `right` stays on `right`). Only swap sides when there is a clear motivated reason (camera reversal, explicit movement across frame, cut/transition), and state that change in `<below>` or dialogue.

Output:
- Return ONLY the final <comic> XML, no extra explanation.
"""
    return prompt


def generate_story_for_dimension(
    openai_client: openai.OpenAI,
    together_client: openai.OpenAI,
    api_model_name: str,
    model_dir: str,
    dimension_id: str,
    story_labels: List[str],
    template_xml: str,
    pose_db: Dict[str, str],
    gen_per_story: int = 2,
) -> None:
    dim_info = DIMENSIONS[dimension_id]
    char_count = dim_info["char_count"]
    story_count_type = dim_info["story_count"]

    for story_label in story_labels:
        story_ids = [int(s) for s in story_label.split("_")]
        if len(story_ids) == 1:
            elements = load_story_elements(story_ids[0])
        else:
            elements = merge_story_elements(story_ids)

        elements_prompt = annotate_elements_for_prompt(elements, pose_db)

        for gen_idx in range(1, gen_per_story + 1):
            print(
                f"  -> Generating: model={model_dir}, dim={dimension_id}, "
                f"story={story_label}, gen{gen_idx}"
            )
            user_prompt = build_user_prompt(
                dimension_id,
                story_label,
                elements_prompt,
                template_xml,
                char_count,
                story_count_type,
            )
            messages = [
                {"role": "system", "content": "You are a strict ComicXML generation assistant."},
                {"role": "user", "content": user_prompt},
            ]

            client = pick_client(api_model_name, openai_client, together_client)
            xml_text = chat_model(client, api_model_name, messages)
            time.sleep(SLEEP_BETWEEN_CALLS)

            out_dir = OUTPUT_BASE / model_dir / dimension_id / f"story_{story_label}"
            out_dir.mkdir(parents=True, exist_ok=True)
            out_path = out_dir / f"gen{gen_idx}.xml"
            out_path.write_text(xml_text, encoding="utf-8")
            print(f"     Saved: {out_path}")


def run_4dim_generation(
    selected_models: List[str],
    story_ids: List[int],
    mode: str,
    pose_db_path: Path,
    gen_per_story: int = 2,
) -> None:
    if not TOGETHER_API_KEY.strip() or not OPENAI_API_KEY.strip():
        print(
            "Please set OPENAI_API_KEY and TOGETHER_API_KEY at the top of this file.",
            file=sys.stderr,
        )
        sys.exit(1)

    if OPENAI_BASE_URL.strip():
        openai_client = openai.OpenAI(
            api_key=OPENAI_API_KEY.strip(), base_url=OPENAI_BASE_URL.strip()
        )
    else:
        openai_client = openai.OpenAI(api_key=OPENAI_API_KEY.strip())

    together_client = openai.OpenAI(
        api_key=TOGETHER_API_KEY.strip(),
        base_url=TOGETHER_BASE_URL.strip() or "https://api.together.xyz/v1",
    )

    if not selected_models:
        selected_models = list(MODELS.keys())

    if not story_ids:
        raise ValueError("You must provide at least one story ID via --stories.")

    if mode not in {"single", "multi"}:
        mode = "single" if len(story_ids) == 1 else "multi"

    story_ids_sorted = sorted(story_ids)
    story_label = "_".join(str(sid) for sid in story_ids_sorted)
    template_xml = load_template()
    pose_db = load_pose_database(pose_db_path)

    print("\n" + "=" * 80)
    print("4-Dimension Constrained Generation (pose DB + panel continuity)")
    print("=" * 80)
    print(f"Output base    : {OUTPUT_BASE}")
    print(f"Pose database  : {pose_db_path} ({len(pose_db)} explicit entries)")
    print(f"Mode           : {mode} (story_ids = {story_ids_sorted}, label = {story_label})")
    print(f"Models to run  : {', '.join(selected_models)}")
    print(f"Generations/Dim: {gen_per_story}")
    print("=" * 80 + "\n")

    if mode == "single":
        if len(story_ids_sorted) != 1:
            raise ValueError("Single-story mode expects exactly ONE story ID.")
        dims_to_run = [
            "dim1_single_char_single_story",
            "dim3_multi_char_single_story",
        ]
    else:
        if len(story_ids_sorted) < 2:
            raise ValueError("Multi-story mode expects at least TWO story IDs.")
        dims_to_run = [
            "dim2_single_char_multi_story",
            "dim4_multi_char_multi_story",
        ]

    for display_name in selected_models:
        if display_name not in MODELS:
            print(f"[WARN] Unknown model name (skipped): {display_name}")
            continue

        model_dir, _prompt_type, api_model_name = MODELS[display_name]

        print("\n" + "=" * 80)
        print(f"Running model: {display_name} (dir={model_dir})")
        print("=" * 80)

        for dim_id in dims_to_run:
            print(f"\n[Dimension] {dim_id}")
            generate_story_for_dimension(
                openai_client=openai_client,
                together_client=together_client,
                api_model_name=api_model_name,
                model_dir=model_dir,
                dimension_id=dim_id,
                story_labels=[story_label],
                template_xml=template_xml,
                pose_db=pose_db,
                gen_per_story=gen_per_story,
            )

    print("\nAll generations completed.")


def run_interactive(pose_db_path: Path) -> None:
    print("\n" + "=" * 80)
    print("Interactive 4-Dimension Constrained Generation (pose DB experiment)")
    print("=" * 80)
    print("Available models:")
    for idx, name in enumerate(MODELS.keys(), start=1):
        mark = " *" if name in FIVE_MODEL_DISPLAY_NAMES else ""
        print(f"  {idx}. {name}{mark}")
    print("  (* = same five as all_eval_4outputs.py)")
    print("  all. Run all models in the list above")
    print(f"  five. Run recommended five: {', '.join(FIVE_MODEL_DISPLAY_NAMES)}")

    model_input = input(
        "\nSelect model(s) by number (e.g., 7 or 1,2), 'five', or 'all' [five]: "
    ).strip()
    if not model_input or model_input.lower() == "five":
        selected_models = list(dict.fromkeys(FIVE_MODEL_DISPLAY_NAMES))
    elif model_input.lower() == "all":
        selected_models = list(MODELS.keys())
    else:
        selected_models = []
        for token in model_input.split(","):
            token = token.strip()
            if not token:
                continue
            if token.isdigit():
                i = int(token)
                if 1 <= i <= len(MODELS):
                    selected_models.append(list(MODELS.keys())[i - 1])
                else:
                    raise ValueError(f"Invalid model index: {i}")
            elif token in MODELS:
                selected_models.append(token)
            else:
                raise ValueError(f"Invalid model selector: {token}")
        selected_models = list(dict.fromkeys(selected_models))

    stories_input = input("Enter story IDs separated by spaces (e.g., '5 6 7'): ").strip()
    if not stories_input:
        raise ValueError("You must provide at least one story ID.")
    story_ids = [int(x) for x in stories_input.split()]

    gen_input = input("Generations per story per dimension [default 2]: ").strip()
    gen_per_story = int(gen_input) if gen_input else 2

    run_4dim_generation(
        selected_models=selected_models,
        story_ids=story_ids,
        mode="",
        pose_db_path=pose_db_path,
        gen_per_story=gen_per_story,
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Constrained 4-dim generation with pose descriptions + panel continuity."
    )
    parser.add_argument(
        "--model",
        type=str,
        nargs="+",
        default=None,
        choices=list(MODELS.keys()),
        help="Models to run. Default in CLI mode: five-model set.",
    )
    parser.add_argument(
        "--stories",
        type=int,
        nargs="+",
        default=None,
        help="Story IDs for element libraries, e.g. --stories 2 3 5.",
    )
    parser.add_argument(
        "--mode",
        type=str,
        default=None,
        choices=["single", "multi"],
        help="single -> dim1+dim3; multi -> dim2+dim4. Omitted: inferred from # of story IDs.",
    )
    parser.add_argument(
        "--gen-per-story",
        type=int,
        default=2,
        help="Generations per story per dimension (default 2).",
    )
    parser.add_argument(
        "--pose-db",
        type=str,
        default=None,
        help=(
            f"Pose database: .xlsx (e.g. Pose/Type/Depicts), .csv, or .json. "
            f"Default: {DEFAULT_POSE_DB}"
        ),
    )
    args = parser.parse_args()

    if args.pose_db:
        pose_db_path = Path(args.pose_db).expanduser()
        if not pose_db_path.is_absolute():
            cwd_try = (Path.cwd() / pose_db_path).resolve()
            exp_try = (EXP_DIR / pose_db_path).resolve()
            root_try = (PROJECT_ROOT / pose_db_path).resolve()
            if cwd_try.exists():
                pose_db_path = cwd_try
            elif exp_try.exists():
                pose_db_path = exp_try
            elif root_try.exists():
                pose_db_path = root_try
            else:
                pose_db_path = cwd_try
        else:
            pose_db_path = pose_db_path.resolve()
    else:
        pose_db_path = DEFAULT_POSE_DB

    if args.stories is None:
        run_interactive(pose_db_path)
        return

    models = args.model
    if models is None:
        models = FIVE_MODEL_DISPLAY_NAMES

    run_4dim_generation(
        selected_models=models,
        story_ids=list(args.stories),
        mode=args.mode or "",
        pose_db_path=pose_db_path,
        gen_per_story=args.gen_per_story,
    )


if __name__ == "__main__":
    if len(sys.argv) == 1:
        pose_path = DEFAULT_POSE_DB
        run_interactive(pose_path)
    else:
        main()
