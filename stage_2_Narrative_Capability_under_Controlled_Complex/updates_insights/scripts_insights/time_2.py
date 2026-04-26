#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import copy
import csv
import json
import re
import sys
import time
from pathlib import Path
from typing import Any, Dict, List

import openai

# CONFIG — paste keys and model ids here
OPENAI_API_KEY = ""
TOGETHER_API_KEY = ""

OPENAI_BASE_URL = ""
TOGETHER_BASE_URL = "https://api.together.xyz/v1"

# API model strings (change if your account uses different slugs)
OPENAI_MODEL_GPT41 = "gpt-4.1"
OPENAI_MODEL_GPT4O = "gpt-4o"
TOGETHER_MODEL_DEEPSEEK_V3 = "deepseek-ai/DeepSeek-V3.1"
TOGETHER_MODEL_DEEPSEEK_R1 = "deepseek-ai/DeepSeek-R1"
TOGETHER_MODEL_LLAMA = "meta-llama/Llama-3.3-70B-Instruct-Turbo"

MAX_TOKENS = 8192
TEMPERATURE = 0.4
SLEEP_BETWEEN_CALLS = 1.0

XML_DECLARATION = '<?xml version="1.0" encoding="UTF-8"?>'

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
POSE_EXP_DIR = PROJECT_ROOT / "constrained_4dim_pose_db"
SPEC_STORY_ELEMENTS_DIR = PROJECT_ROOT / "spec_library"
TEMPLATE_FILE = PROJECT_ROOT / "template.xml"
_XLSX_DEFAULT = PROJECT_ROOT / "pose database.xlsx"
DEFAULT_POSE_DB = _XLSX_DEFAULT if _XLSX_DEFAULT.exists() else (POSE_EXP_DIR / "pose_database.json")

OUTPUT_BASE = PROJECT_ROOT / "outputs_insights" / "time_2"
DIMENSION_ID = "dim4_multi_char_multi_story"

DEFAULT_POSE_FALLBACK_DESCRIPTION = (
    "No separate gloss in pose_database; interpret staging from the canonical pose name."
)

TOGETHER_API_MODEL_IDS = frozenset(
    {
        TOGETHER_MODEL_DEEPSEEK_V3,
        TOGETHER_MODEL_DEEPSEEK_R1,
        TOGETHER_MODEL_LLAMA,
    }
)


def load_pose_database_xlsx(path: Path) -> Dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError("Reading .xlsx requires openpyxl. pip install openpyxl") from e

    mapping: Dict[str, str] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(min_row=1, values_only=True)
        header = next(rows, None)
        if not header:
            return {}
        names = [(str(c).strip() if c is not None else "").lower() for c in header]

        def col_idx(*candidates: str) -> int | None:
            for i, n in enumerate(names):
                if n in candidates:
                    return i
            return None

        i_pose = col_idx("pose")
        if i_pose is None:
            raise ValueError(f"Excel {path} needs a 'Pose' column; got headers {list(header)}")
        i_type = col_idx("type")
        i_depicts = col_idx("depicts", "description", "desc")

        def cell(row: tuple, idx: int | None) -> str:
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            if v is None:
                return ""
            return str(v).strip()

        for row in rows:
            if row is None:
                continue
            pose = cell(row, i_pose)
            if not pose:
                continue
            parts: List[str] = []
            dep = cell(row, i_depicts)
            typ = cell(row, i_type)
            if dep:
                parts.append(dep)
            if typ:
                parts.append(f"type: {typ}")
            mapping[pose] = ". ".join(parts) if parts else ""
    finally:
        wb.close()
    return mapping


def load_pose_database(path: Path) -> Dict[str, str]:
    if not path.exists():
        return {}
    suf = path.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        return load_pose_database_xlsx(path)
    if suf == ".csv":
        mapping: Dict[str, str] = {}
        with path.open("r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                return {}
            pkey = None
            dkey = None
            for h in reader.fieldnames:
                hl = h.strip().lower()
                if hl in ("pose", "pose_name", "name", "id"):
                    pkey = h
                if hl in ("description", "desc", "gloss", "text"):
                    dkey = h
            if not pkey or not dkey:
                raise ValueError(
                    f"CSV {path} needs pose+description columns; got {reader.fieldnames}"
                )
            for row in reader:
                k = (row.get(pkey) or "").strip()
                v = (row.get(dkey) or "").strip()
                if k:
                    mapping[k] = v
        return mapping
    if suf != ".json":
        raise ValueError(f"Unsupported pose database format: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict) and all(isinstance(k, str) for k in data):
        raw = data["poses"] if "poses" in data and isinstance(data["poses"], dict) else data
        return {str(k): str(v) if v is not None else "" for k, v in raw.items()}
    raise ValueError(f"Unsupported pose database format: {path}")


def annotate_elements_for_prompt(elements: Dict[str, Any], pose_db: Dict[str, str]) -> Dict[str, Any]:
    out = copy.deepcopy(elements)
    poses = out.get("poses", [])
    if not isinstance(poses, list):
        return out
    new_poses: List[str] = []
    for p in poses:
        if not isinstance(p, str):
            continue
        desc = (pose_db.get(p) or "").strip()
        if not desc:
            desc = DEFAULT_POSE_FALLBACK_DESCRIPTION
        new_poses.append(f"{p} ({desc})")
    out["poses"] = new_poses
    return out


def load_template() -> str:
    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError(f"template.xml not found at {TEMPLATE_FILE}")
    return TEMPLATE_FILE.read_text(encoding="utf-8")


def load_story_elements(story_id: int) -> Dict:
    story_file = SPEC_STORY_ELEMENTS_DIR / f"story_{story_id}.json"
    if not story_file.exists():
        raise FileNotFoundError(f"Story elements file not found: {story_file}")
    with story_file.open("r", encoding="utf-8") as f:
        return json.load(f)


def merge_story_elements(story_ids: List[int]) -> Dict:
    merged = {
        "characters": {},
        "poses": set(),
        "settings": set(),
        "facings": set(),
        "positions": set(),
        "balloon_status": set(),
    }
    for sid in story_ids:
        data = load_story_elements(sid)
        for char_id, char_data in data.get("characters", {}).items():
            if char_id not in merged["characters"]:
                merged["characters"][char_id] = char_data
            else:
                existing_variants = merged["characters"][char_id].get("variants", {})
                new_variants = char_data.get("variants", {})
                existing_variants.update(new_variants)
                merged["characters"][char_id]["variants"] = existing_variants
        merged["poses"].update(data.get("poses", []))
        merged["settings"].update(data.get("settings", []))
        merged["facings"].update(data.get("facings", []))
        merged["positions"].update(data.get("positions", []))
        merged["balloon_status"].update(data.get("balloon_status", []))
    return {
        "characters": merged["characters"],
        "poses": sorted(merged["poses"]),
        "settings": sorted(merged["settings"]),
        "facings": sorted(merged["facings"]),
        "positions": sorted(merged["positions"]),
        "balloon_status": sorted(merged["balloon_status"]),
    }


def build_user_prompt(
    dimension_id: str,
    story_label: str,
    elements_for_prompt: Dict,
    template_xml: str,
    char_count: int,
    story_count: str,
) -> str:
    dim_desc = {
        "dim1_single_char_single_story": "Dim1: Single character and single story.",
        "dim2_single_char_multi_story": "Dim2: Single character and multiple stories merged.",
        "dim3_multi_char_single_story": "Dim3: Multiple characters within a single story.",
        "dim4_multi_char_multi_story": "Dim4: Multiple characters and multiple stories merged.",
    }.get(dimension_id, dimension_id)

    if char_count == 1:
        char_constraint = (
            "- You MUST use only ONE main character ID from the element library\n"
            "  across the whole story (single-character setting).\n"
        )
    else:
        char_constraint = (
            "- You MUST use MULTIPLE distinct character IDs from the element library\n"
            "  across the story (multi-character setting), and most panels should\n"
            "  contain 2 characters.\n"
        )

    story_info = (
        f"You must create a ComicXML story for **{dim_desc}**\n"
        f"- Story source type: {story_count} (story_label = {story_label})\n"
        "- Panel count must be between 6 and 12 (inclusive).\n"
        f"{char_constraint}"
    )

    elements_json = json.dumps(elements_for_prompt, ensure_ascii=False, indent=2)

    prompt = f"""
You are an expert ComicXML generation assistant.

Your task:
- Generate a COMPLETE ComicXML for a comic story.
- STRICTLY follow the structure defined by the template.
- Use ONLY elements from the provided element library.
- Respect the character constraints above (single-character or multi-character).

{story_info}

Element library (permitted values only). Poses are listed as **canonical_pose_name (description)**.
The parenthetical text is ONLY a staging gloss for you; in the final XML you must still obey
the canonical element rules below.

{elements_json}

ComicXML template (STRUCTURE GUIDE, not actual values):
{template_xml}

IMPORTANT CONSTRAINTS:
1. XML must be well-formed (proper tag closure and nesting).
2. XML must follow the template structure (no missing required elements).
3. Use only elements (characters, poses, settings, positions, facings, balloon_status) from
   the element library above. For **poses**, each entry is `CANONICAL_NAME (description)`.
   The canonical name may contain spaces; it is the substring **before the first ` (`**
   (space + opening parenthesis). Put **only** that canonical name inside `<pose>...</pose>`,
   never the parenthetical description.
4. You MAY reuse elements across panels, but do not invent new ones.
5. Panel count must be between 6 and 12.
6. For multi-character dimensions, most panels should contain multiple characters.
7. Pay close attention to facing directions in multi-character panels.
8. If two characters are in a face-to-face conversation across left and right positions, the usual orientation is:
   - left-position character: facing="right"
   - right-position character: facing="left"
9. In non-conversational or non-face-to-face situations, choose facing directions according to the narrative context of that panel.
10. **Cast visibility across panels (multi-character dimensions):**
   - Default: If panel N and panel N+1 depict the **same ongoing scene** and the **same interaction or conversation** continues, keep the **same characters visibly on stage** in panel N+1 unless the story motivates a change.
   - Do **not** drop a co-present character between adjacent panels **without a clear diegetic reason** (e.g. they exit on-panel, the `<below>` caption states a time jump or location change, or the dialogue explicitly addresses their absence).
   - A **single-character panel** is allowed when: (a) the story **explicitly** moves to a new beat with only one character; (b) a `<below>` caption states a transition; or (c) the plot states the other character left or is off-screen for a named reason.
   - After **physical actions involving two characters** (rescuing, dousing, grabbing, face-to-face dialogue), the **next panel should normally still show both**, repositioned if needed, rather than removing one character for a solo middle composition without motivation.
11. **Layout:** When two characters are **co-present in the same narrative beat**, prefer **left+right** (or **left+middle+right**) over a **solo `<middle>`** panel. Do not use solo middle only to emphasize a speech balloon if both characters should still be in the scene.
12. **`<below>` captions:** For panels where **multiple characters share a beat** (especially immediately after joint action or dialogue), write each relevant `<below>` so it **mentions where each of them is** (stance or position relative to the other), not only one of them — unless a transition or solo beat is intentional and stated in that caption.
13. **Global story logic and continuity:** The whole comic must read as one coherent story arc (setup -> development -> consequence/resolution), not disconnected moments. Maintain clear cause-effect links across adjacent panels. If the story changes location, time, or cast composition, explicitly signal that change in `<below>` and/or dialogue.
14. **Element whitelist (strict):** Do not output any out-of-library value, including near-synonyms or morphology variants (e.g. `laughing` when only `reacting` exists). If an intended action is unavailable, pick the closest allowed element from the provided library.
15. **Cross-panel position consistency (multi-character dimensions):** If the same two characters continue the same interaction across adjacent panels in the same scene, keep their screen sides stable by default (the character on `left` stays on `left`, `right` stays on `right`). Only swap sides when there is a clear motivated reason (camera reversal, explicit movement across frame, cut/transition), and state that change in `<below>` or dialogue.

Output:
- Return ONLY the final <comic> XML, no extra explanation.
"""
    return prompt


def inject_speech_constraints(prompt: str) -> str:
    extra = (
        "16. **Speech coverage (default):** Unless the scene is explicitly and naturally non-verbal "
        "(e.g., establishing shot, silent transition, purely physical action beat), include speech content "
        "in the panel.\n"
        "17. **One speech per panel:** Each panel should contain at most ONE `<balloon status=\"speech\">` "
        "with one `<content>`.\n"
        "18. **Dialogue distribution:** If a conversation is longer, spread it across multiple panels "
        "rather than adding multiple speech balloons in one panel.\n"
        "19. **`<below>` brevity:** Keep each `<below>` to ONE concise sentence. Do not add speculative "
        "details that are not supported by the panel context.\n"
    )
    marker = "\nOutput:\n"
    if marker in prompt:
        return prompt.replace(marker, f"{extra}\nOutput:\n", 1)
    return (
        f"{prompt.rstrip()}\n\n{extra}\nOutput:\n"
        "- Return ONLY the final <comic> XML, no extra explanation.\n"
    )


def strip_markdown_fences(text: str) -> str:
    s = text.strip()
    if not s.startswith("```"):
        return s
    lines = s.splitlines()
    if lines and lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip().startswith("```"):
        lines = lines[:-1]
    return "\n".join(lines).strip()


def extract_comic_xml_file_body(raw: str) -> str:
    """
    Keep from optional <?xml ...?> through </comic>.
    If the model omits the XML declaration, prepend a standard UTF-8 declaration.
    """
    s = strip_markdown_fences(raw)
    low = s.lower()
    end_idx = low.rfind("</comic>")
    if end_idx < 0:
        raise ValueError("Model output does not contain a closing </comic> tag.")
    end_idx += len("</comic>")

    xml_decl_match = re.search(r"<\?xml\b[^>]*\?>", s, flags=re.IGNORECASE)
    comic_start = low.find("<comic")
    if comic_start < 0:
        raise ValueError("Model output does not contain <comic>.")

    if xml_decl_match and xml_decl_match.start() <= comic_start:
        start = xml_decl_match.start()
        fragment = s[start:end_idx].strip()
        return fragment

    comic_fragment = s[comic_start:end_idx].strip()
    return f"{XML_DECLARATION}\n{comic_fragment}"


def chat_completion(client: openai.OpenAI, model: str, messages: list, max_tokens: int, temperature: float) -> str:
    params: Dict[str, Any] = {"model": model, "messages": messages, "temperature": temperature}
    if model.startswith("o1"):
        params["max_completion_tokens"] = max_tokens
    else:
        params["max_tokens"] = max_tokens
    try:
        resp = client.chat.completions.create(**params)
        choice = resp.choices[0]
        content = (choice.message.content or "").strip()
        fr = getattr(choice, "finish_reason", None)
        if fr == "length":
            print(
                f"[WARN] Completion hit max_tokens limit (finish_reason=length, cap={max_tokens}). "
                "Increase MAX_TOKENS or shorten the prompt."
            )
        if content:
            return content
        raise RuntimeError("API returned empty content.")
    except Exception as e:
        raise RuntimeError(f"API call failed: {e}") from e


def pick_client(
    api_model: str,
    openai_client: openai.OpenAI,
    together_client: openai.OpenAI,
) -> openai.OpenAI:
    if api_model in TOGETHER_API_MODEL_IDS:
        return together_client
    return openai_client


def main() -> None:
    if not TOGETHER_API_KEY.strip() or not OPENAI_API_KEY.strip():
        print(
            "Please set OPENAI_API_KEY and TOGETHER_API_KEY at the top of this file "
            "(run_dim4_speech_openai_together.py).",
            file=sys.stderr,
        )
        sys.exit(1)

    if OPENAI_BASE_URL.strip():
        openai_client = openai.OpenAI(
            api_key=OPENAI_API_KEY.strip(), base_url=OPENAI_BASE_URL.strip()
        )
    else:
        openai_client = openai.OpenAI(api_key=OPENAI_API_KEY.strip())

    together_client = openai.OpenAI(
        api_key=TOGETHER_API_KEY.strip(),
        base_url=TOGETHER_BASE_URL.strip() or "https://api.together.xyz/v1",
    )

    print("Dim4 only | speech constraints | five models: 2×OpenAI + 3×Together\n")

    stories_raw = input("Enter story source IDs (space-separated, at least 2, e.g. 5 6 7): ").strip()
    if not stories_raw:
        print("No story IDs given.", file=sys.stderr)
        sys.exit(1)
    story_ids = [int(x) for x in stories_raw.split()]
    if len(story_ids) < 2:
        print("Dim4 requires at least two story IDs.", file=sys.stderr)
        sys.exit(1)
    story_ids_sorted = sorted(story_ids)
    story_label = "_".join(str(s) for s in story_ids_sorted)

    gen_raw = input("Generations per model [2]: ").strip()
    gen_per_model = int(gen_raw) if gen_raw else 2
    if gen_per_model < 1:
        print("Generations must be >= 1.", file=sys.stderr)
        sys.exit(1)

    pose_db = load_pose_database(DEFAULT_POSE_DB)
    template_xml = load_template()
    elements = merge_story_elements(story_ids_sorted)
    elements_prompt = annotate_elements_for_prompt(elements, pose_db)

    char_count = 2
    story_count = "multi"

    models: List[tuple[str, str, str]] = [
        ("gpt-4.1", "gpt_4", OPENAI_MODEL_GPT41),
        ("gpt-4o", "gpt_4o", OPENAI_MODEL_GPT4O),
        ("DeepSeek-V3", "deepseek_v3", TOGETHER_MODEL_DEEPSEEK_V3),
        ("DeepSeek-R1", "deepseek_r1", TOGETHER_MODEL_DEEPSEEK_R1),
        ("llama-3.3-70B", "llama_3_70b", TOGETHER_MODEL_LLAMA),
    ]

    base_prompt = build_user_prompt(
        DIMENSION_ID,
        story_label,
        elements_prompt,
        template_xml,
        char_count,
        story_count,
    )
    user_prompt = inject_speech_constraints(base_prompt)
    messages = [
        {"role": "system", "content": "You are a strict ComicXML generation assistant."},
        {"role": "user", "content": user_prompt},
    ]

    print(f"\nOutput base: {OUTPUT_BASE}")
    print(f"Pose DB: {DEFAULT_POSE_DB} ({len(pose_db)} entries)")
    print(f"Story label: {story_label} | gens per model: {gen_per_model}\n")

    for display_name, model_dir, api_model in models:
        client = pick_client(api_model, openai_client, together_client)
        backend = "Together" if api_model in TOGETHER_API_MODEL_IDS else "OpenAI"
        print("=" * 80)
        print(f"{display_name} ({backend}) api_model={api_model}")
        print("=" * 80)
        for gen_idx in range(1, gen_per_model + 1):
            print(f"  -> gen {gen_idx}/{gen_per_model} ...")
            try:
                raw = chat_completion(client, api_model, messages, MAX_TOKENS, TEMPERATURE)
            except RuntimeError as e:
                out_err = (
                    OUTPUT_BASE
                    / model_dir
                    / DIMENSION_ID
                    / f"story_{story_label}"
                    / f"gen{gen_idx}_error.txt"
                )
                out_err.parent.mkdir(parents=True, exist_ok=True)
                out_err.write_text(str(e) + "\n", encoding="utf-8")
                print(f"     [ERROR] {e} — error saved to {out_err}")
                # Continue to next generation/model instead of aborting the whole run.
                continue
            try:
                xml_only = extract_comic_xml_file_body(raw)
            except ValueError as e:
                out_debug = (
                    OUTPUT_BASE
                    / model_dir
                    / DIMENSION_ID
                    / f"story_{story_label}"
                    / f"gen{gen_idx}_raw.txt"
                )
                out_debug.parent.mkdir(parents=True, exist_ok=True)
                out_debug.write_text(raw, encoding="utf-8")
                print(f"     [ERROR] {e} — raw saved to {out_debug}")
                continue

            out_path = (
                OUTPUT_BASE
                / model_dir
                / DIMENSION_ID
                / f"story_{story_label}"
                / f"gen{gen_idx}.xml"
            )
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_text(xml_only + "\n", encoding="utf-8")
            print(f"     Saved: {out_path}")
            time.sleep(SLEEP_BETWEEN_CALLS)

    print("\nAll models finished.")


if __name__ == "__main__":
    main()
